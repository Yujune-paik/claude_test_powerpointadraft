import os
import tempfile
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pptx import Presentation

app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def _extract_slide_text(slide) -> str:
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    texts.append(text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        texts.append(text)
    return "\n".join(texts)


@app.post("/api/upload")
async def upload_pptx(file: UploadFile = File(...)):
    if not file.filename.endswith((".pptx", ".PPTX")):
        raise HTTPException(status_code=400, detail="Only .pptx files are supported")

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        content = await file.read()
        tmp.write(content)
        tmp_path = tmp.name

    try:
        prs = Presentation(tmp_path)
        slides_data = []
        for i, slide in enumerate(prs.slides):
            slides_data.append({
                "index": i,
                "text": _extract_slide_text(slide),
            })

        return {
            "sessionId": "vercel-session",
            "slideCount": len(slides_data),
            "slides": slides_data,
        }
    finally:
        os.unlink(tmp_path)

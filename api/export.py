import os
import tempfile
import base64
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pptx import Presentation
import json

app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.post("/api/export")
async def export_pptx(
    file: UploadFile = File(...),
    notes: str = Form(...),
):
    """Receive original PPTX + notes JSON, return PPTX with notes embedded."""
    notes_dict = json.loads(notes)

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        content = await file.read()
        tmp.write(content)
        tmp_path = tmp.name

    try:
        prs = Presentation(tmp_path)

        for slide_index_str, note_text in notes_dict.items():
            idx = int(slide_index_str)
            if idx < len(prs.slides):
                slide = prs.slides[idx]
                if not slide.has_notes_slide:
                    slide.notes_slide
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = note_text

        output_path = tmp_path + "_with_notes.pptx"
        prs.save(output_path)

        with open(output_path, "rb") as f:
            pptx_bytes = f.read()

        os.unlink(output_path)

        return Response(
            content=pptx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": "attachment; filename=presentation_with_notes.pptx"},
        )
    finally:
        os.unlink(tmp_path)
